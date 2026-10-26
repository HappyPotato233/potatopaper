# 问答路由：多 Agent 编排问答（ask/steps/result）+ 方向多视角综述（对齐详细设计 §3.5）
import io
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from ai import orchestrator
from app.deps import get_current_user
from app.models import Subscription
from app.schemas.qa import (AskRequest, AskResponse, ReviewExportRequest,
                            ReviewRequest, StepsResponse)
from app.session import get_db

router = APIRouter(prefix="/api/qa", tags=["qa"])


@router.post("/ask", response_model=AskResponse, status_code=202)
def ask(body: AskRequest, user: dict = Depends(get_current_user)):
    """发起提问：返回会话 id；服务端后台启动 查询改写+检索路由 → v4 主-子编排（或 v3 回退）"""
    session_id = orchestrator.ask(user["id"], body.query, mode=body.mode,
                                  rewrite=body.rewrite, enable_planner=body.enable_planner)
    return AskResponse(session_id=session_id)


@router.get("/{session_id}/steps", response_model=StepsResponse)
def steps(session_id: str, user: dict = Depends(get_current_user)):
    """轮询多步状态（8 节点：查询改写→规划调度→…→复审校验→汇总答案，含编排字段）"""
    data = orchestrator.get_steps(session_id)
    if not data:
        raise HTTPException(status_code=404, detail="会话不存在或已过期")
    return StepsResponse(**data)


@router.get("/{session_id}/result")
def result(session_id: str, user: dict = Depends(get_current_user)):
    """拉取最终回答（含引用来源）；未完成时 finished=False"""
    data = orchestrator.get_result(session_id)
    if not data:
        raise HTTPException(status_code=404, detail="会话不存在或已过期")
    return data


@router.post("/review", status_code=202)
def review(body: ReviewRequest, user: dict = Depends(get_current_user),
           session: Session = Depends(get_db)):
    """多视角综述（异步）：校验订阅归属后创建综述会话，前端轮询 steps/result"""
    sub = session.scalar(select(Subscription.id).where(
        Subscription.id == body.subscription_id, Subscription.user_id == user["id"]))
    if sub is None:
        raise HTTPException(status_code=404, detail="订阅不存在")
    session_id = orchestrator.ask_review(user["id"], body.subscription_id)
    return {"session_id": session_id, "status": "queued"}


@router.post("/review/export")
def review_export(body: ReviewExportRequest, user: dict = Depends(get_current_user)):
    """综述导出 Word（F-REV-05）：接收用户编辑后的大纲 → 返回 .docx 二进制流（不落库）"""
    data = _build_review_docx(body.title, body.outline, body.citations)
    filename = quote(f"{body.title or '文献综述'}.docx")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


def _build_review_docx(title: str, outline: list, citations: list) -> bytes:
    """按大纲层级生成综述文档：标题层级 + 正文分段 + 本章引用 + 参考文献列表

    依赖 python-docx，延迟导入：缺依赖时仅本接口报错，不影响其他链路。
    """
    from docx import Document

    doc = Document()
    doc.add_heading(title or "文献综述", level=0)
    for node in outline:
        doc.add_heading(node.title or "（未命名章节）", level=min(max(node.depth, 1), 3))
        for para in (node.content or "").split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        if node.citations:
            doc.add_paragraph("本章引用：" + " ".join(f"[{n}]" for n in node.citations))
    doc.add_heading("参考文献", level=1)
    if not citations:
        doc.add_paragraph("（无参考文献）")
    for i, c in enumerate(citations, start=1):
        parts = " | ".join(x for x in (c.title, c.source, c.year, c.url) if x)
        doc.add_paragraph(f"[{i}] {parts}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
